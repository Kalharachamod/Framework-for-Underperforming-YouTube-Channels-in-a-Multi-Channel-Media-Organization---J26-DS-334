"""Experiment 3C: researcher annotation batch 1 from the expanded development pool.

Actionable demand is rare (6 of 40 calibration comments, of which half were
rule-selected), so one sampling source cannot both find enough positives and
measure how common demand is. Batch 1 therefore mixes, in one blinded and
shuffled workbook:

- ``existing_queue``: the 94 Experiment 2A.2 and 60 Experiment 2B review rows;
- ``random``: a channel-stratified random sample of the development pool, the
  only rows that may be used to estimate prevalence or what retrieval misses;
- ``retrieval_rule`` / ``retrieval_tfidf`` / ``retrieval_e5``: candidates from
  the frozen rule baseline, combined TF-IDF and multilingual-e5 similarity to
  the six researcher-approved calibration demands. Comparing their verified
  yield is itself a result.

Only ``development_pool`` and ``expanded_development_pool`` comments are
eligible (``ExpandedCorpusGuard``). Held-out, quarantined and calibration rows
are never read as candidates. Provenance lives in a private key file; the
workbook shows no source, score or rule output.
"""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd

from .annotation_schema import ACTIONABILITY_LABELS, PRIMARY_SIGNALS, SPAM_STATUS_LABELS
from .baseline import detect_frame
from .build_expanded_corpus import DEFAULT_OUTPUT_DIR as CORPUS_DIR
from .build_expanded_corpus import ExpandedCorpusGuard, TRAINING_ROLES
from .prepare_experiment_2a2 import DEFAULT_CALIBRATION
from .prepare_experiment_2b import build_tfidf_representations, load_calibration_query_seeds
from .supply_topics import E5_MODEL, E5_REVISION, build_supply_documents


SEED = 20261021
RANDOM_PER_CHANNEL = 25
RETRIEVAL_PER_METHOD = 50
DESCRIPTION_CHARS = 400
PACKAGE_DIR = Path(__file__).resolve().parent
OUTPUTS = PACKAGE_DIR / "outputs"
DEFAULT_QUEUES = (
    OUTPUTS / "experiment_2a2" / "development_candidate_control_review.csv",
    OUTPUTS / "experiment_2b" / "tfidf_novel_candidate_review.csv",
)
DEFAULT_OUTPUT_DIR = OUTPUTS / "experiment_3c"
RETRIEVAL_METHODS = ("retrieval_rule", "retrieval_tfidf", "retrieval_e5")

# Visible annotation columns use the schema v1 annotator_1 names so the
# completed workbook imports without renaming.
ANNOTATION_COLUMNS = [
    "annotator_1_primary_signal",
    "annotator_1_actionable_content_demand",
    "annotator_1_requested_action",
    "annotator_1_requested_topic",
    "annotator_1_request_target",
    "annotator_1_spam_status",
    "annotator_1_notes",
]
CONTEXT_COLUMNS = [
    "review_id",
    "channel_name",
    "video_title",
    "video_context",
    "video_published",
    "is_reply",
    "parent_comment_text",
    "comment_published",
    "comment_text",
    "video_url",
]


def _rank(value: str, namespace: str, seed: int = SEED) -> str:
    return hashlib.sha256(f"{seed}:{namespace}:{value}".encode("utf-8")).hexdigest()


def load_development(corpus_dir: Path = CORPUS_DIR) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Return (all comments for parent lookup, guarded development comments)."""

    comments = pd.read_csv(corpus_dir / "expanded_comments.csv", dtype=str, keep_default_na=False)
    guard = ExpandedCorpusGuard.load(corpus_dir / "comment_partition_manifest_v2.csv")
    development = comments.loc[comments["dataset_role"].isin(TRAINING_ROLES)].copy()
    guard.validate_training_frame(development[["comment_id"]])
    return comments, development.reset_index(drop=True)


def stratified_random(development: pd.DataFrame, excluded: set[str], per_channel: int) -> pd.DataFrame:
    eligible = development.loc[~development["comment_id"].isin(excluded)].copy()
    eligible["rank"] = eligible["comment_id"].map(lambda v: _rank(v, "random"))
    picked = eligible.sort_values(["channel_name", "rank"]).groupby("channel_name").head(per_channel)
    return picked.assign(source="random").drop(columns="rank")


def _round_robin(ranked: list[list[str]], limit: int, excluded: set[str]) -> list[str]:
    chosen: list[str] = []
    seen = set(excluded)
    position = 0
    while len(chosen) < limit and any(position < len(r) for r in ranked):
        for candidates in ranked:
            if position < len(candidates) and candidates[position] not in seen:
                chosen.append(candidates[position])
                seen.add(candidates[position])
                if len(chosen) >= limit:
                    break
        position += 1
    return chosen


def retrieval_pool(development: pd.DataFrame) -> pd.DataFrame:
    """Rows worth retrieving: not emoji-only and with some text."""

    keep = development["is_emoji_only"].str.lower().ne("true") & development["cleaned_text"].str.strip().ne("")
    return development.loc[keep].reset_index(drop=True)


def rule_candidates(pool: pd.DataFrame, excluded: set[str], limit: int) -> tuple[list[str], int]:
    frame = pool.rename(columns={"raw_comment_text": "original_text"})
    flags = detect_frame(frame[["original_text", "cleaned_text"]])["candidate_for_content_demand"].astype(bool)
    flagged = pool.loc[flags.to_numpy()].copy()
    flagged["rank"] = flagged["comment_id"].map(lambda v: _rank(v, "rule"))
    per_channel = [
        group.sort_values("rank")["comment_id"].tolist()
        for _, group in flagged.groupby("channel_name", sort=True)
    ]
    return _round_robin(per_channel, limit, excluded), int(flags.sum())


def similarity_candidates(scores: np.ndarray, ids: list[str], excluded: set[str], limit: int) -> list[str]:
    """Round-robin over seeds, each seed's neighbours by descending positive similarity."""

    ranked = []
    for seed_scores in scores:
        order = np.argsort(-seed_scores, kind="stable")
        ranked.append([ids[i] for i in order if seed_scores[i] > 0])
    return _round_robin(ranked, limit, excluded)


def e5_scores(corpus_texts: list[str], query_texts: list[str]) -> np.ndarray:
    import torch
    from sentence_transformers import SentenceTransformer

    torch.manual_seed(0)
    model = SentenceTransformer(E5_MODEL, revision=E5_REVISION, device="cpu")
    encode = lambda texts: model.encode(  # noqa: E731
        [f"query: {t}" for t in texts], batch_size=64, normalize_embeddings=True, show_progress_bar=False
    )
    return encode(query_texts) @ encode(corpus_texts).T


def video_context(corpus_dir: Path) -> pd.DataFrame:
    videos = pd.read_csv(corpus_dir / "expanded_videos.csv", dtype=str, keep_default_na=False)
    documents = build_supply_documents(videos)
    documents["video_context"] = [
        doc[len(title) + 3 :][:DESCRIPTION_CHARS] if doc.startswith(title + " . ") else ""
        for doc, title in zip(documents["document"], documents["title"])
    ]
    return documents.set_index("video_id")[["title", "video_context", "published_at"]]


def build_workbook_rows(selected: pd.DataFrame, comments: pd.DataFrame, videos: pd.DataFrame) -> pd.DataFrame:
    text_by_id = comments.set_index("comment_id")["raw_comment_text"]
    rows = selected.copy()
    rows["shuffle"] = rows["comment_id"].map(lambda v: _rank(v, "order"))
    rows = rows.sort_values("shuffle").reset_index(drop=True)
    rows["review_id"] = [f"B1-{i + 1:04d}" for i in range(len(rows))]
    rows["video_title"] = rows["video_id"].map(videos["title"]).fillna("")
    rows["video_context"] = rows["video_id"].map(videos["video_context"]).fillna("")
    rows["video_published"] = rows["video_id"].map(videos["published_at"]).fillna("").str[:10]
    rows["parent_comment_text"] = rows["parent_comment_id"].map(text_by_id).fillna("")
    rows.loc[rows["is_reply"].str.lower().eq("true") & rows["parent_comment_text"].eq(""), "parent_comment_text"] = (
        "(reply; parent comment not collected)"
    )
    rows["comment_published"] = rows["published_at"].str[:10]
    rows["comment_text"] = rows["raw_comment_text"]
    rows["video_url"] = "https://www.youtube.com/watch?v=" + rows["video_id"]
    for column in ANNOTATION_COLUMNS:
        rows[column] = ""
    return rows


def write_workbook(rows: pd.DataFrame, path: Path, guide_lines: list[str] | None = None) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    guide = wb.active
    guide.title = "Guide"
    for line in guide_lines or GUIDE_LINES:
        guide.append([line])
    guide.column_dimensions["A"].width = 140
    for cell in guide["A"]:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    guide["A1"].font = Font(bold=True, size=14)

    sheet = wb.create_sheet("Annotate")
    headers = CONTEXT_COLUMNS + ANNOTATION_COLUMNS
    sheet.append(headers)
    for record in rows[headers].itertuples(index=False):
        sheet.append(list(record))
    widths = {
        "review_id": 10, "channel_name": 16, "video_title": 40, "video_context": 40, "video_published": 12,
        "is_reply": 8, "parent_comment_text": 30, "comment_published": 12, "comment_text": 60, "video_url": 18,
        "annotator_1_primary_signal": 30, "annotator_1_actionable_content_demand": 14,
        "annotator_1_requested_action": 30, "annotator_1_requested_topic": 26, "annotator_1_request_target": 18,
        "annotator_1_spam_status": 14, "annotator_1_notes": 36,
    }
    header_fill = PatternFill("solid", fgColor="DDE7F5")
    label_fill = PatternFill("solid", fgColor="FFF4CC")
    for index, header in enumerate(headers, start=1):
        letter = sheet.cell(row=1, column=index).column_letter
        sheet.column_dimensions[letter].width = widths.get(header, 14)
        sheet.cell(row=1, column=index).font = Font(bold=True)
        sheet.cell(row=1, column=index).fill = label_fill if header in ANNOTATION_COLUMNS else header_fill
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.freeze_panes = "B2"
    last = len(rows) + 1
    for header, values in (
        ("annotator_1_primary_signal", PRIMARY_SIGNALS),
        ("annotator_1_actionable_content_demand", ACTIONABILITY_LABELS),
        ("annotator_1_spam_status", SPAM_STATUS_LABELS),
    ):
        letter = sheet.cell(row=1, column=headers.index(header) + 1).column_letter
        validation = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
        validation.error = "Choose a value from the list."
        sheet.add_data_validation(validation)
        validation.add(f"{letter}2:{letter}{last}")
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


GUIDE_LINES = [
    "Component 2 annotation batch 1 — how to annotate",
    "",
    "Work in the 'Annotate' sheet. Rows are shuffled, so any stretch of rows is a fair mix: you can stop and resume anywhere.",
    "For each row read the comment with its video title, video context and (for replies) the parent comment. Then fill the yellow columns.",
    "",
    "1. annotator_1_primary_signal (one value, dropdown):",
    "   EXPLICIT_CONTENT_REQUEST — asks for a new video, topic, guest, interview, re-upload, full episode or coverage.",
    "   CONTENT_CHANGE_PREFERENCE — wants an existing programme changed: storyline, cast, judges, format, length, language, subtitles.",
    "   INFORMATION_NEED — asks for an explanation, details, clarification or a follow-up.",
    "   CONTENT_QUALITY_FEEDBACK — specific production feedback (audio, camera, editing, ads).",
    "   TOPICAL_INTEREST_OR_DISCUSSION — talks about the topic without asking the channel to produce or change anything.",
    "   ORDINARY_REACTION — praise, criticism, thanks, emojis, wishes, support for a contestant.",
    "   SPAM_OR_IRRELEVANT — promotion, scams, unrelated self-promotion, bot-like text.",
    "   UNCLEAR — not enough context to decide. Do not guess the author's intention.",
    "",
    "2. annotator_1_actionable_content_demand (YES / NO / UNCLEAR):",
    "   YES only if the channel could act on it: a concrete request or change the channel itself could make.",
    "   A request aimed at someone else (a politician, a contestant's family) is NO. Praise with 'please' in it is NO.",
    "   Implicit requests count (e.g. 'we don't feel like watching without him' = bring him back).",
    "",
    "3. If YES: annotator_1_requested_action (what should be done) and annotator_1_requested_topic (about what).",
    "   annotator_1_request_target: the programme or person the request is about, if clear.",
    "",
    "4. annotator_1_spam_status: SPAM / NOT_SPAM / UNCLEAR. Negative or rude comments are not spam.",
    "5. annotator_1_notes: anything ambiguous (missing parent, sarcasm, mixed signals).",
    "",
    "Keep the review_id and comment_text unchanged. Leave a row blank if you skip it; blank means 'not yet annotated'.",
]


def prepare(
    *,
    corpus_dir: Path = CORPUS_DIR,
    queue_paths: tuple[Path, ...] = DEFAULT_QUEUES,
    calibration_path: Path = DEFAULT_CALIBRATION,
    output_dir: Path = DEFAULT_OUTPUT_DIR,
    random_per_channel: int = RANDOM_PER_CHANNEL,
    retrieval_per_method: int = RETRIEVAL_PER_METHOD,
    use_e5: bool = True,
) -> dict[str, Any]:
    comments, development = load_development(corpus_dir)
    by_id = development.set_index("comment_id", drop=False)

    calibration = pd.read_csv(calibration_path, dtype=str, keep_default_na=False)
    seeds = load_calibration_query_seeds(calibration_path)
    queue_ids: list[str] = []
    for path in queue_paths:
        queue_ids += pd.read_csv(path, dtype=str, keep_default_na=False)["comment_id"].tolist()
    missing = sorted(set(queue_ids).difference(by_id.index))
    if missing:
        raise ValueError(f"{len(missing)} existing queue comments are outside the guarded development pool")
    taken = set(queue_ids) | set(calibration["comment_id"])

    parts = [by_id.loc[list(dict.fromkeys(queue_ids))].assign(source="existing_queue")]
    random_rows = stratified_random(development, taken, random_per_channel)
    parts.append(random_rows)
    taken |= set(random_rows["comment_id"])

    pool = retrieval_pool(development)
    pool_ids = pool["comment_id"].tolist()
    rule_ids, rule_flagged = rule_candidates(pool, taken, retrieval_per_method)
    taken |= set(rule_ids)
    corpus_matrices, query_matrices, _ = build_tfidf_representations(pool["normalized_text"], seeds["original_text"])
    tfidf = (query_matrices["combined"] @ corpus_matrices["combined"].T).toarray()
    tfidf_ids = similarity_candidates(tfidf, pool_ids, taken, retrieval_per_method)
    taken |= set(tfidf_ids)
    e5_ids: list[str] = []
    if use_e5:
        e5 = e5_scores(pool["normalized_text"].tolist(), seeds["original_text"].tolist())
        e5_ids = similarity_candidates(e5, pool_ids, taken, retrieval_per_method)
    for name, ids in (("retrieval_rule", rule_ids), ("retrieval_tfidf", tfidf_ids), ("retrieval_e5", e5_ids)):
        parts.append(by_id.loc[ids].assign(source=name))

    selected = pd.concat(parts, ignore_index=True)
    if selected["comment_id"].duplicated().any():
        raise ValueError("A comment was selected by more than one source")
    ExpandedCorpusGuard.load(corpus_dir / "comment_partition_manifest_v2.csv").validate_training_frame(
        selected[["comment_id"]]
    )

    rows = build_workbook_rows(selected, comments, video_context(corpus_dir))
    output_dir.mkdir(parents=True, exist_ok=True)
    write_workbook(rows, output_dir / "annotation_batch1.xlsx")
    rows[CONTEXT_COLUMNS + ANNOTATION_COLUMNS].to_csv(output_dir / "annotation_batch1.csv", index=False, encoding="utf-8-sig")
    rows[["review_id", "comment_id", "video_id", "channel_name", "dataset_role", "source"]].to_csv(
        output_dir / "annotation_batch1_private_key.csv", index=False, encoding="utf-8"
    )
    summary = {
        "rows": int(len(rows)),
        "by_source": rows["source"].value_counts().to_dict(),
        "by_channel": rows["channel_name"].value_counts().to_dict(),
        "by_source_channel": pd.crosstab(rows["source"], rows["channel_name"]).to_dict(orient="index"),
        "development_pool_comments": int(len(development)),
        "retrieval_pool_comments": int(len(pool)),
        "rule_flagged_in_retrieval_pool": rule_flagged,
        "seeds": int(len(seeds)),
        "seed": SEED,
        "e5": {"model": E5_MODEL, "revision": E5_REVISION} if use_e5 else None,
    }
    (output_dir / "annotation_batch1_summary.json").write_text(json.dumps(summary, indent=2) + "\n", encoding="utf-8")
    return summary


def main() -> int:
    parser = argparse.ArgumentParser(description="Prepare Experiment 3C annotation batch 1.")
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--no-e5", action="store_true", help="Skip the multilingual-e5 retrieval method.")
    args = parser.parse_args()
    print(json.dumps(prepare(output_dir=args.output_dir, use_e5=not args.no_e5), indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
