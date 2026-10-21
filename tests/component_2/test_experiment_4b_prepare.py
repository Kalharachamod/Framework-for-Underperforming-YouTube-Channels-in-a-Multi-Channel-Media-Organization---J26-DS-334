from __future__ import annotations

import tempfile
import unittest
from collections import Counter
from pathlib import Path

import pandas as pd

from research.component_2.feasibility_experiment.prepare_experiment_4b import (
    ANNOTATION_COLUMNS,
    CONTEXT_COLUMNS,
    MAX_PER_VIDEO,
    contestant_table,
    draw_samples,
    guide_lines,
    latin_variants,
    mention_mask,
    sinhala_variants,
    title_performers,
    write_validation_workbook,
)


class TitleParsingTests(unittest.TestCase):
    def test_reads_performer_after_song_title(self):
        self.assertEqual(
            title_performers("Seetha Raaka (සීත රෑක මධුසමේ) Pahanmi kandhanarachchige | Dream Star Season 12 | TV Derana"),
            ["Pahanmi Kandhanarachchige"],
        )
        self.assertEqual(title_performers("Rahase (රහසේ) Nipuni Anarkalee Ft Ravi Royster | Dream Star"), ["Nipuni Anarkalee"])

    def test_group_songs_and_titles_without_performer(self):
        self.assertEqual(title_performers("Mal Mal Sariya (මල් මල් සාරිය) Group Song | Dream Star Season 12"), [])
        self.assertEqual(title_performers("🔴 LIVE | Derana Dream Star Season 12 | Top 08"), [])
        self.assertEqual(title_performers("ගායනයෙන් හද සොරාගත් ජනනි... | Dream Star Season 12"), [])

    def test_contestant_table_merges_first_name_and_marks_main(self):
        videos = pd.DataFrame({
            "video_id": ["a", "b", "c", "d"],
            "title": [
                "S (ස) Pahanmi Kandhanarachchige | DS", "T (ත) Pahanmi | DS",
                "U (උ) Pahanmi Kandhanarachchige | DS", "V (ව) Venuri Nirmanika | DS",
            ],
            "published_at": ["2025-04-02T00:00:00Z", "2025-05-01T00:00:00Z", "2025-06-01T00:00:00Z", "2025-04-03T00:00:00Z"],
        })
        table = contestant_table(videos, min_main=3).set_index("first_name")
        self.assertEqual(table.loc["Pahanmi", "videos"], 3)
        self.assertTrue(table.loc["Pahanmi", "is_main"])
        self.assertFalse(table.loc["Venuri", "is_main"])
        self.assertEqual(table.loc["Pahanmi", "last_upload"], "2025-06-01")


class SpellingTests(unittest.TestCase):
    def test_latin_variants_keep_similar_frequent_tokens(self):
        counts = Counter({"pahanmi": 50, "pahanmita": 9, "pahnmi": 4, "pahana": 2, "hello": 40, "pahanmige": 1})
        found = dict(latin_variants("Pahanmi", counts, min_count=3))
        self.assertEqual(set(found), {"pahanmi", "pahanmita", "pahnmi"})

    def test_sinhala_variants_use_stems(self):
        counts = Counter({"පහන්මි": 20, "පහන්මිට": 5, "පහන": 30, "පහම්මි": 4})
        found = dict(sinhala_variants(("පහන්ම", "පහම්"), counts, min_count=3))
        self.assertEqual(set(found), {"පහන්මි", "පහන්මිට", "පහම්මි"})

    def test_mention_mask_matches_whole_latin_tokens_and_sinhala_prefixes(self):
        texts = pd.Series(["Go Pahanmi!", "pahanmita denna", "පහන්මිගෙ සින්දු", "pahana", "nothing here"])
        self.assertEqual(mention_mask(texts, ["pahanmi", "පහන්මි"]).tolist(), [True, False, True, False, False])


def _meaningful(n_videos: int = 30, per_video: int = 10) -> pd.DataFrame:
    rows = []
    for v in range(n_videos):
        for i in range(per_video):
            kind = i % 3
            text = ["pahanmi super", "judges are unfair", "lassanai"][kind]
            rows.append({"comment_id": f"c{v}_{i}", "video_id": f"v{v}", "normalized_text": f"{text} {v}"})
    return pd.DataFrame(rows)


class SamplingTests(unittest.TestCase):
    def test_strata_exclusions_overlap_and_video_cap(self):
        meaningful = _meaningful()
        excluded = {"c0_0", "c1_0"}
        strata = {"contestant_mention": 30, "judges_or_voting": 20, "random": 20}
        sample, prompt_dev = draw_samples(meaningful, ["pahanmi"], excluded, strata=strata, prompt_dev_size=10)
        self.assertEqual(sample["stratum"].value_counts().to_dict(), strata)
        self.assertEqual(len(prompt_dev), 10)
        self.assertFalse(set(sample["comment_id"]) & set(prompt_dev["comment_id"]))
        self.assertFalse(set(sample["comment_id"]) & excluded)
        self.assertTrue(sample["comment_id"].is_unique)
        self.assertLessEqual(sample["video_id"].value_counts().max(), MAX_PER_VIDEO)
        mention = sample.loc[sample["stratum"].eq("contestant_mention"), "normalized_text"]
        self.assertTrue(mention.str.contains("pahanmi").all())
        judges = sample.loc[sample["stratum"].eq("judges_or_voting"), "normalized_text"]
        self.assertTrue(judges.str.contains("judges").all())

    def test_sampling_is_deterministic(self):
        meaningful = _meaningful()
        strata = {"contestant_mention": 10, "judges_or_voting": 5, "random": 5}
        first, _ = draw_samples(meaningful, ["pahanmi"], set(), strata=strata, prompt_dev_size=5)
        second, _ = draw_samples(meaningful.sample(frac=1, random_state=3), ["pahanmi"], set(), strata=strata, prompt_dev_size=5)
        self.assertEqual(sorted(first["comment_id"]), sorted(second["comment_id"]))


class WorkbookTests(unittest.TestCase):
    def test_blind_workbook_has_no_model_or_sampling_columns(self):
        from openpyxl import load_workbook

        rows = pd.DataFrame([{c: "" for c in CONTEXT_COLUMNS + ANNOTATION_COLUMNS} for _ in range(3)])
        rows["review_id"] = ["V4B-001", "V4B-002", "V4B-003"]
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "sheet.xlsx"
            write_validation_workbook(rows, ["Janani", "Pahanmi"], path)
            wb = load_workbook(path)
            self.assertEqual(wb.sheetnames, ["Guide", "Annotate", "Lists"])
            self.assertEqual(wb["Lists"].sheet_state, "hidden")
            headers = [cell.value for cell in wb["Annotate"][1]]
            self.assertEqual(headers, CONTEXT_COLUMNS + ANNOTATION_COLUMNS)
            self.assertNotIn("stratum", headers)
            self.assertNotIn("comment_id", headers)
            self.assertEqual(len(wb["Annotate"].data_validations.dataValidation), 7)
            targets = [wb["Lists"].cell(row=r, column=1).value for r in range(1, 9)]
            self.assertEqual(targets[:3], ["Janani", "Pahanmi", "OTHER_CONTESTANT"])
        self.assertTrue(any("Do not use AI" in line for line in guide_lines(["Janani"])))


if __name__ == "__main__":
    unittest.main()
